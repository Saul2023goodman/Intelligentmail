import tempfile
import unittest
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from smartmail import SmartMail, SmartMailError
from smartmail.mailbox import ControlledMailbox
from tests import test_execution
from tests.test_duplicates import outbound_sent_observation


def master(path, rows, headers=None, merges=()):
    """Anonymized fixture using the inspected master list's layout."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or ["大学", "导师", "邮箱📮", "URL"])
    for row in rows:
        sheet.append(row)
    for cells in merges:
        sheet.merge_cells(cells)
    workbook.save(path)
    workbook.close()
    return path


def revised_master(path, rows, note, headers=None):
    """Same supported rows in a workbook with different bytes (an extra column)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or ["大学", "导师", "邮箱📮", "URL", "备注"])
    for row in rows:
        sheet.append([*row, note])
    workbook.save(path)
    workbook.close()
    return path


class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.home = self.directory / "state"
        self.core = SmartMail(self.home)
        self.addCleanup(self.core.__exit__)
        self.campaign = self.core.create_campaign("2027 outreach")

    def test_import_establishes_inspectable_task_and_participants(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["Example University", "Dr Alex Green", "alex@example.edu", "https://example.edu/people/alex"],
        ])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        task = self.core.get_task(result["task_ids"][0])
        self.assertEqual(task["student"]["name"], "Test Student")
        self.assertEqual(task["mailbox"]["address"], "student@163.com")
        self.assertEqual(task["supervisor"]["name"], "Dr Alex Green")
        self.assertEqual(task["supervisor"]["addresses"], ["alex@example.edu"])
        self.assertEqual(task["institution"]["name"], "Example University")
        self.assertEqual(task["campaign"], self.campaign)
        self.assertEqual(self.core.list_tasks(self.campaign["id"])[0]["id"], task["id"])
        self.assertEqual(task["source_associations"][0]["row"], 2)
        self.assertEqual(task["exceptions"], [])

    def test_reordered_alias_columns_and_merged_institutions_keep_cell_evidence(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            [" a@example.edu\n", "Example University", "Dr Alex Green"],
            ["b@example.edu", None, "Dr Blair Blue"],
        ], headers=["邮箱", "院校", "导师姓名"], merges=["B2:B3"])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        tasks = [self.core.get_task(i) for i in result["task_ids"]]
        self.assertEqual([t["institution"]["name"] for t in tasks], ["Example University"] * 2)
        self.assertEqual(tasks[0]["supervisor"]["addresses"], ["a@example.edu"])
        evidence = tasks[1]["source_associations"][0]["evidence"]
        self.assertEqual(evidence["field_cells"]["institution"], "B2")
        self.assertEqual(evidence["raw_cells"]["B3"], None)

    def test_bundle_sources_and_tasks_survive_restart_and_original_removal(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [["University", "Alex", "a@example.edu", ""]])
        bundle = self.directory / "sample.zip"
        with zipfile.ZipFile(bundle, "w") as archive:
            archive.write(source, "学生/名单.xlsx")
            archive.writestr("学生/CV.docx", b"original document bytes")
        original = source.read_bytes()
        result = self.core.import_master(self.campaign["id"], student["id"], bundle)
        source.unlink()
        bundle.unlink()
        with SmartMail(self.home) as restarted:
            imported = restarted.get_import(result["id"])
            self.assertEqual(len(imported["sources"]), 3)
            workbook = next(s for s in imported["sources"] if s["name"].endswith(".xlsx"))
            self.assertEqual(restarted.read_source(workbook["id"]), original)
            opened = restarted.materialize_source(workbook["id"])
            self.assertEqual(opened.read_bytes(), original)
            opened.write_bytes(b"operator edits the opened copy")
            self.assertEqual(restarted.materialize_source(workbook["id"]).read_bytes(), original)
            self.assertEqual(restarted.get_task(result["task_ids"][0])["campaign"], self.campaign)

    def test_profile_url_in_email_column_becomes_a_persisted_blocker(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [[
            "University", "Prof Alex", "https://example.edu/people/alex/#", "https://example.edu/people/alex/",
        ]])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        with SmartMail(self.home) as restarted:
            task = restarted.get_task(result["task_ids"][0])
            self.assertEqual(task["supervisor"]["addresses"], [])
            self.assertEqual(task["exceptions"][0]["code"], "invalid_recipient")
            self.assertTrue(task["exceptions"][0]["blocking"])
            self.assertEqual(task["source_associations"][0]["evidence"]["raw_cells"]["C2"], "https://example.edu/people/alex/#")

    def test_same_profile_evidence_shares_supervisor_and_task_across_known_addresses(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["University", "Dr Alex Green", "a@example.edu", "https://example.edu/people/alex/"],
            ["University", "Alex Green", "alternate@example.edu", "https://example.edu/people/alex/#"],
        ])
        first = self.core.import_master(self.campaign["id"], student["id"], source)
        second = self.core.import_master(self.campaign["id"], student["id"], source)
        self.assertEqual(first["task_ids"], second["task_ids"])
        self.assertEqual(len(first["task_ids"]), 1)
        self.assertTrue(second["duplicate"])
        self.assertEqual(second["id"], first["id"])
        self.assertEqual({row["outcome"] for row in second["rows"]}, {"reused"})
        self.assertEqual(second["summary"]["new_sources"], 0)
        task = self.core.get_task(first["task_ids"][0])
        self.assertEqual(task["supervisor"]["addresses"], ["a@example.edu", "alternate@example.edu"])
        self.assertEqual(len(task["source_associations"]), 2)
        self.assertEqual(len(self.core.list_imports(self.campaign["id"])), 1)
        other_campaign = self.core.create_campaign("2028 outreach")
        other = self.core.import_master(other_campaign["id"], student["id"], source)
        self.assertNotEqual(first["task_ids"], other["task_ids"])
        self.assertEqual(self.core.get_task(other["task_ids"][0])["supervisor"]["id"], task["supervisor"]["id"])
        other_student = self.core.create_student("Another Student", "another@163.com")
        separate = self.core.import_master(self.campaign["id"], other_student["id"], source)
        self.assertNotEqual(first["task_ids"], separate["task_ids"])

    def test_unreliable_identity_matches_remain_separate_and_block_both_tasks(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["University", "Alex Green", "a@example.edu", ""],
            ["University", "Alex Green", "b@example.edu", ""],
        ])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        with SmartMail(self.home) as restarted:
            tasks = [restarted.get_task(i) for i in result["task_ids"]]
            self.assertEqual(len({t["supervisor"]["id"] for t in tasks}), 2)
            for task in tasks:
                self.assertIn("identity_ambiguity", [e["code"] for e in task["exceptions"]])
                self.assertEqual(len(task["supervisor"]["addresses"]), 1)

    def test_shared_addresses_do_not_override_conflicting_profile_evidence(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["University", "Alex Green", "shared@example.edu", "https://example.edu/people/alex-one"],
            ["University", "Alex Green", "shared@example.edu", "https://example.edu/people/alex-two"],
            ["Other University", "Blair Blue", "shared@example.edu", "https://other.edu/blair"],
        ])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        self.assertEqual(len(result["task_ids"]), 3)
        for task_id in result["task_ids"]:
            self.assertIn("identity_ambiguity", [e["code"] for e in self.core.get_task(task_id)["exceptions"]])

    def test_unsupported_or_ambiguous_layouts_reject_without_partial_import(self):
        student = self.core.create_student("Test Student", "student@163.com")
        for headers, rows in [
            (["大学", "导师", "邮箱", "Email"], [["University", "Alex", "a@example.edu", "b@example.edu"]]),
            (["大学", "导师", "邮箱"], [["University", "Alex", "a@example.edu"], [None, "Blair", "b@example.edu"]]),
            (["大学", "导师", "邮箱"], [["University", "=A2", "a@example.edu"]]),
        ]:
            with self.subTest(headers=headers, rows=rows):
                source = master(self.directory / "master.xlsx", rows, headers=headers)
                with self.assertRaises(SmartMailError):
                    self.core.import_master(self.campaign["id"], student["id"], source)
                self.assertEqual(self.core.list_tasks(self.campaign["id"]), [])
                self.assertEqual(self.core.list_imports(self.campaign["id"]), [])

    def test_reimport_of_the_same_unresolved_source_row_keeps_task_identity(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [["University", "Alex Green", "", ""]])
        first = self.core.import_master(self.campaign["id"], student["id"], source)
        second = self.core.import_master(self.campaign["id"], student["id"], source)
        self.assertEqual(first["task_ids"], second["task_ids"])
        self.assertEqual(len(self.core.list_tasks(self.campaign["id"])), 1)
        task = self.core.get_task(first["task_ids"][0])
        self.assertEqual({e["code"] for e in task["exceptions"]}, {"invalid_recipient"})

    def test_new_profile_evidence_is_retained_for_later_address_matching(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["University", "Alex Green", "a@example.edu", ""],
            ["University", "Alex Green", "a@example.edu", "https://example.edu/alex"],
            ["University", "Alex Green", "b@example.edu", "https://example.edu/alex"],
        ])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        self.assertEqual(len(result["task_ids"]), 1)
        self.assertEqual(self.core.get_task(result["task_ids"][0])["exceptions"], [])

    def test_campaign_selection_and_task_summary_expose_unfinished_work(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [["University", "Alex Green", "", ""]])
        for campaign_id, student_id in [(None, student["id"]), (self.campaign["id"], "unknown")]:
            with self.assertRaises(SmartMailError):
                self.core.import_master(campaign_id, student_id, source)
        self.assertEqual(self.core.list_imports(self.campaign["id"]), [])
        result = self.core.import_master(self.campaign["id"], student["id"], source)
        summary = self.core.list_tasks(self.campaign["id"])[0]
        self.assertEqual(summary["id"], result["task_ids"][0])
        self.assertEqual(summary["student_name"], "Test Student")
        self.assertEqual(summary["supervisor_name"], "Alex Green")
        self.assertEqual(summary["institution_name"], "University")
        self.assertEqual(summary["exception_count"], 1)
        empty = self.core.create_campaign("Separate campaign")
        self.assertEqual(self.core.list_tasks(empty["id"]), [])


class DuplicateImportTests(IntakeTests):
    def test_exact_duplicate_import_is_idempotent_and_reported_per_row(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["Example University", "Dr Alex Green", "alex@example.edu", ""],
        ])
        first = self.core.import_master(self.campaign["id"], student["id"], source)
        second = self.core.import_master(self.campaign["id"], student["id"], source)

        self.assertTrue(second["duplicate"])
        self.assertEqual(second["id"], first["id"])
        self.assertEqual(second["summary"], {
            "rows": 1, "new": 0, "reused": 1, "duplicate": 0,
            "conflicts": 0, "new_sources": 0,
        })
        self.assertEqual(len(self.core.list_imports(self.campaign["id"])), 1)
        self.assertEqual(len(self.core.get_import(first["id"])["sources"]), 1)
        task = self.core.get_task(first["task_ids"][0])
        self.assertEqual(len(task["source_associations"]), 1)

    def test_a_revision_bundle_preserves_new_materials_but_reuses_master_rows(self):
        student = self.core.create_student("Test Student", "student@163.com")
        master_path = master(self.directory / "master.xlsx", [
            ["Example University", "Dr Alex Green", "alex@example.edu", ""],
        ])
        first_bundle = self.directory / "bundle-1.zip"
        with zipfile.ZipFile(first_bundle, "w") as archive:
            archive.write(master_path, "master.xlsx")
            archive.writestr("CV.docx", b"original document bytes")
        first = self.core.import_master(self.campaign["id"], student["id"], first_bundle)

        second_bundle = self.directory / "bundle-2.zip"
        with zipfile.ZipFile(second_bundle, "w") as archive:
            archive.write(master_path, "master.xlsx")
            archive.writestr("Example University_Dr Alex Green.docx", b"revised letter bytes")
        second = self.core.import_master(self.campaign["id"], student["id"], second_bundle)

        self.assertFalse(second["duplicate"])
        self.assertNotEqual(second["id"], first["id"])
        self.assertEqual(second["summary"]["new_sources"], 2)
        self.assertEqual([source["name"] for source in second["reused_sources"]], ["master.xlsx"])
        self.assertEqual(
            sorted(source["name"] for source in self.core.get_import(second["id"])["sources"]),
            ["Example University_Dr Alex Green.docx", "bundle-2.zip"])
        self.assertEqual(second["rows"][0]["outcome"], "reused")
        self.assertEqual(len(self.core.list_imports(self.campaign["id"])), 2)
        task = self.core.get_task(first["task_ids"][0])
        self.assertEqual(len(task["source_associations"]), 1)
        self.assertEqual(len(self.core.get_import(first["id"])["sources"]), 3)

    def test_duplicate_rows_within_one_master_keep_one_task_and_a_nonblocking_exception(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["Example University", "Dr Alex Green", "alex@example.edu", ""],
            ["Example University", "Dr Alex Green", "alex@example.edu", ""],
        ])
        result = self.core.import_master(self.campaign["id"], student["id"], source)

        self.assertEqual(len(result["task_ids"]), 1)
        self.assertEqual([row["outcome"] for row in result["rows"]], ["new", "duplicate"])
        self.assertEqual(result["summary"]["duplicate"], 1)
        task = self.core.get_task(result["task_ids"][0])
        self.assertEqual(len(task["source_associations"]), 2)
        duplicate = next(exception for exception in task["exceptions"]
                         if exception["code"] == "duplicate_import_row")
        self.assertFalse(duplicate["blocking"])
        self.assertIn("Sheet!3", duplicate["detail"])
        self.assertIn("Sheet!2", duplicate["detail"])

    def test_rows_adding_evidence_are_reused_not_duplicates(self):
        student = self.core.create_student("Test Student", "student@163.com")
        source = master(self.directory / "master.xlsx", [
            ["University", "Dr Alex Green", "a@example.edu", ""],
            ["University", "Dr Alex Green", "a@example.edu", "https://example.edu/alex"],
            ["University", "Dr Alex Green", "b@example.edu", "https://example.edu/alex"],
        ])
        result = self.core.import_master(self.campaign["id"], student["id"], source)

        self.assertEqual(len(result["task_ids"]), 1)
        self.assertEqual([row["outcome"] for row in result["rows"]],
                         ["new", "reused", "reused"])
        changes = {change["code"] for row in result["rows"] for change in row["changes"]}
        self.assertEqual(changes, {"profile_added", "address_added"})
        task = self.core.get_task(result["task_ids"][0])
        self.assertEqual(task["supervisor"]["addresses"],
                         ["a@example.edu", "b@example.edu"])
        self.assertNotIn("duplicate_import_row",
                         [exception["code"] for exception in task["exceptions"]])

    def test_reimport_with_changed_information_reports_changes_without_blocking(self):
        student = self.core.create_student("Test Student", "student@163.com")
        rows = [["University", "Dr Alex Green", "a@example.edu", ""]]
        self.core.import_master(
            self.campaign["id"], student["id"],
            master(self.directory / "master-1.xlsx", rows))
        with_profile = self.core.import_master(
            self.campaign["id"], student["id"],
            revised_master(self.directory / "master-2.xlsx",
                           [["University", "Dr Alex Green", "a@example.edu",
                             "https://example.edu/alex"]], "v2"))
        self.assertEqual(with_profile["rows"][0]["outcome"], "reused")
        self.assertIn("profile_added",
                      [change["code"] for change in with_profile["rows"][0]["changes"]])

        with_alternate = self.core.import_master(
            self.campaign["id"], student["id"],
            revised_master(self.directory / "master-3.xlsx",
                           [["University", "Dr Alex Green", "b@example.edu",
                             "https://example.edu/alex"]], "v3"))
        self.assertIn("address_added",
                      [change["code"] for change in with_alternate["rows"][0]["changes"]])

        without_address = self.core.import_master(
            self.campaign["id"], student["id"],
            revised_master(self.directory / "master-4.xlsx",
                           [["University", "Dr Alex Green", "",
                             "https://example.edu/alex"]], "v4"))
        self.assertEqual(without_address["rows"][0]["outcome"], "reused")
        self.assertIn("address_absent_in_row",
                      [change["code"] for change in without_address["rows"][0]["changes"]])

        task = self.core.get_task(with_profile["task_ids"][0])
        self.assertEqual(task["supervisor"]["addresses"],
                         ["a@example.edu", "b@example.edu"])
        self.assertEqual(task["exceptions"], [])


class ImportPriorOutreachConflictTests(test_execution.ExecutionTestCase):
    def observe(self, observation):
        self.core.mailbox = ControlledMailbox(observations=[observation])
        return self.core.refresh_mailbox(self.student["id"])

    def import_letter(self):
        (cv_name, cv_path), _ = self.cv()
        return self.import_bundle(
            [(test_execution.DRAFT_NAME,
              test_execution.draft_paragraphs(
                  "alex@example.edu", "Dear Dr Green,", [test_execution.DECLARATION]))],
            extra=[(cv_name, cv_path)])

    def test_observed_prior_send_blocks_imported_initial_outreach_until_resolved(self):
        self.observe(outbound_sent_observation("alex@example.edu"))
        imported = self.import_letter()

        self.assertTrue(imported["rows"][0]["conflict"])
        self.assertEqual(imported["summary"]["conflicts"], 1)
        task = self.core.get_task(imported["task_ids"][0])
        conflict = next(exception for exception in task["exceptions"]
                        if exception["code"] == "prior_outreach_conflict")
        self.assertTrue(conflict["blocking"])
        self.assertIn("outbound sent message", conflict["detail"])

        preparation_id = self.core.prepare_from_documents(imported["id"])["preparation_ids"][0]
        self.core.set_subject(preparation_id, test_execution.SUBJECT)
        slot = self.core.get_preparation(preparation_id)["attachment_slots"][0]
        self.core.confirm_attachment(preparation_id, slot["id"])
        preparation = self.core.get_preparation(preparation_id)
        self.assertFalse(preparation["ready"])
        self.assertIn("prior_outreach_conflict",
                      [finding["code"] for finding in preparation["readiness_findings"]])
        with self.assertRaises(SmartMailError):
            self.core.confirm(preparation_id)

        resolved = self.core.resolve_prior_outreach(task["id"])
        self.assertNotIn("prior_outreach_conflict",
                         [exception["code"] for exception in resolved["exceptions"]])
        self.assertTrue(self.core.get_preparation(preparation_id)["ready"])
        with self.assertRaises(SmartMailError):
            self.core.resolve_prior_outreach(task["id"])

    def test_ambiguous_and_unrelated_observations_do_not_conflict_import(self):
        ambiguous = outbound_sent_observation("alex@example.edu")
        ambiguous["messages"][0]["ambiguity"] = "Recipient success could not be established"
        self.observe(ambiguous)
        first = self.import_bundle([])
        self.assertFalse(first["rows"][0]["conflict"])

        self.observe(outbound_sent_observation("someone-else@example.edu", reference="other-1"))
        second = self.core.import_master(
            self.campaign["id"], self.student["id"],
            revised_master(self.directory / "master-2.xlsx",
                           test_execution.DEFAULT_ROWS, "v2"))
        self.assertFalse(second["rows"][0]["conflict"])
        self.assertEqual(second["rows"][0]["outcome"], "reused")
        self.assertEqual(second["task_ids"], first["task_ids"])

        self.observe(outbound_sent_observation("alex@example.edu", reference="prior-send-2"))
        third = self.core.import_master(
            self.campaign["id"], self.student["id"],
            revised_master(self.directory / "master-3.xlsx",
                           test_execution.DEFAULT_ROWS, "v3"))
        self.assertTrue(third["rows"][0]["conflict"])

    def test_same_campaign_sent_record_conflicts_reimport_but_follow_up_stays_ready(self):
        self.at(datetime(2026, 9, 14, 9, 0, tzinfo=timezone.utc))
        preparation, _ = self.ready_preparation()
        confirmation = self.core.confirm(preparation["id"])
        self.core.run_execution([confirmation["id"]])

        reimported = self.core.import_master(
            self.campaign["id"], self.student["id"],
            revised_master(self.directory / "master-revised.xlsx",
                           test_execution.DEFAULT_ROWS, "revised"))
        self.assertEqual(reimported["rows"][0]["outcome"], "reused")
        self.assertTrue(reimported["rows"][0]["conflict"])
        task = self.core.get_task(preparation["task_id"])
        self.assertIn("prior_outreach_conflict",
                      [exception["code"] for exception in task["exceptions"]])

        self.core.configure_follow_up_rule(
            self.campaign["id"], delay_days=1, maximum_count=1,
            subject_template="Re: {original_subject}",
            body_template="Dear {supervisor_name},\n\nFollowing up.\n\n{student_name}")
        self.at(datetime(2026, 9, 15, 10, 0, tzinfo=timezone.utc))
        actions = self.core.prepare_follow_ups(self.campaign["id"])
        self.assertEqual(len(actions), 1)
        follow_up = self.core.get_preparation(actions[0]["preparation_id"])
        self.assertEqual(follow_up["action_kind"], "follow_up")
        self.assertTrue(follow_up["ready"])
        self.assertEqual(
            self.core.check_duplicate(follow_up["id"])["finding"], "linked_follow_up")

    def test_prior_send_in_another_campaign_is_not_a_conflict_until_observed(self):
        preparation, _ = self.ready_preparation()
        confirmation = self.core.confirm(preparation["id"])
        self.core.run_execution([confirmation["id"]])

        second_campaign = self.core.create_campaign("2028 outreach")
        imported = self.core.import_master(
            second_campaign["id"], self.student["id"],
            revised_master(self.directory / "master-2028.xlsx",
                           test_execution.DEFAULT_ROWS, "2028"))
        self.assertEqual(imported["rows"][0]["outcome"], "new")
        self.assertFalse(imported["rows"][0]["conflict"])

        self.observe(outbound_sent_observation("alex@example.edu"))
        again = self.core.import_master(
            second_campaign["id"], self.student["id"],
            revised_master(self.directory / "master-2028-b.xlsx",
                           test_execution.DEFAULT_ROWS, "2028b"))
        self.assertTrue(again["rows"][0]["conflict"])
        self.assertTrue(any(
            exception["code"] == "prior_outreach_conflict" and exception["blocking"]
            for exception in self.core.get_task(imported["task_ids"][0])["exceptions"]))
