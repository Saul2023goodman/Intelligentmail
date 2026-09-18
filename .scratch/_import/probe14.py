"""Probe the ticket-14 intake recognition gains."""

import io
import tempfile
from pathlib import Path

from openpyxl import Workbook

from smartmail.recognition import recognize_bytes, recognize_file, profile_column, detect_delimiter

TMP = Path(tempfile.mkdtemp())


def wb(name, builder):
    path = TMP / name
    book = Workbook()
    builder(book)
    book.save(path)
    book.close()
    return path


def show(label, result):
    print(f"\n=== {label}")
    print("  type      :", result["type"], result["confidence"])
    for reason in result["reasons"]:
        print("  +", reason)
    for caution in result["cautions"]:
        print("  !", caution)
    print("  identities:", result["identities"])


# 1. institution merged across its supervisors
def merged(book):
    ws = book.active
    ws.title = "名单"
    ws.append(["大学", "导师", "邮箱"])
    ws.append(["A大学", "张三", "a@x.edu"])
    ws.merge_cells("A3:A5")
    ws["A3"] = "B大学"
    ws["B3"], ws["C3"] = "李四", "b@x.edu"
    ws["B4"], ws["C4"] = "王五", "c@x.edu"
    ws["B5"], ws["C5"] = "赵六", "d@x.edu"


show("merged institution roster", recognize_file(wb("merged.xlsx", merged)))


# 2. decorated + English headers
def decorated(book):
    ws = book.active
    ws.append(["Professor", "Affiliation", "E-mail Address", "Homepage"])
    for i in range(3):
        ws.append([f"Prof {i}", f"University {i}", f"p{i}@x.edu", f"https://x.edu/{i}"])


show("decorated/English headers", recognize_file(wb("decorated.xlsx", decorated)))


# 3. header pushed down by a banner and blanks
def deep(book):
    ws = book.active
    ws.append(["2027 导师联系总表"])
    ws.append([])
    ws.append([])
    ws.append(["备注：仅第一轮"])
    ws.append(["大学", "导师", "邮箱", "研究方向"])
    for i in range(3):
        ws.append([f"大学{i}", f"导师{i}", f"t{i}@x.edu", "computational linguistics"])


show("deep header row", recognize_file(wb("deep.xlsx", deep)))


# 4. institution + address, no supervisor column
def partial(book):
    ws = book.active
    ws.append(["院校", "联系邮箱", "主页"])
    for i in range(3):
        ws.append([f"院校{i}", f"g{i}@x.edu", f"https://x.edu/{i}"])


show("institution+address roster", recognize_file(wb("partial.xlsx", partial)))


# 5. semicolon-delimited CSV
semi = "编号;收件人;主题;正文;定时时间\n" + "\n".join(
    f"{i};s{i}@x.edu;Subject {i};Dear Professor,\n\nA complete letter body of "
    f"sufficient length for the row to look like a real message.\n\nYours sincerely,\nStudent;"
    f"2026-09-{10+i} 09:30" for i in range(1, 4))
show("semicolon csv", recognize_bytes("batch.csv", semi.encode("utf-8")))

# 6. tab separated
tsv = "收件人\t主题\t正文\n" + "\n".join(
    f"s{i}@x.edu\tSubject {i}\tDear Professor,\n\nA complete letter body of sufficient "
    f"length for the row to look like a real message.\n\nYours sincerely,\nStudent"
    for i in range(1, 4))
show("tsv", recognize_bytes("batch.tsv", tsv.encode("utf-8")))

# 7. plain text that is not a table
show("prose txt", recognize_bytes("notes.txt", "第一封邮件\n第二封邮件\n".encode("utf-8")))

# 8. delimiter sniffer
print("\n=== detect_delimiter")
print("  comma  :", repr(detect_delimiter('a,b,c\n"Smith, John",x,y')))
print("  semi   :", repr(detect_delimiter('a;b;c\nd;e;f')))
print("  tab    :", repr(detect_delimiter("a\tb\tc\nd\te\tf")))

# 9. profile_column
profile = profile_column(0, "邮箱", ["a@x.edu", "b@x.edu", "not-an-email"])
print("\n=== profile_column", profile)
