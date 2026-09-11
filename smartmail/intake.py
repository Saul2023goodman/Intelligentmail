"""Deterministic extraction for the inspected supervisor master workbook."""

from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

from openpyxl import load_workbook


ALIASES = {
    "institution": {"大学", "学校", "院校", "institution", "university"},
    "name": {"导师", "导师姓名", "supervisor"},
    "address": {"邮箱📮", "邮箱", "电子邮箱", "email"},
    "profile": {"url", "导师主页", "profile url"},
}


class IntakeError(ValueError):
    """The source does not fit the documented intake pattern."""


def read_bundle(path: Path) -> tuple[list[tuple[str, bytes]], int]:
    data = path.read_bytes()
    sources = [(path.name, data)]
    if path.suffix.lower() == ".xlsx":
        return sources, 0
    if path.suffix.lower() != ".zip":
        raise IntakeError("Supported inputs are .xlsx master lists or .zip bundles")
    # The representative Windows archive uses GBK for non-UTF-8 member names.
    with ZipFile(BytesIO(data), metadata_encoding="gbk") as archive:
        for item in archive.infolist():
            name = item.filename.replace("\\", "/")
            if item.is_dir():
                continue
            if PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts or ":" in name:
                raise IntakeError(f"Unsafe archive member name: {name}")
            sources.append((name, archive.read(item)))
    masters = [i for i, (name, _) in enumerate(sources[1:], 1) if name.lower().endswith(".xlsx")]
    if len(masters) != 1:
        raise IntakeError("ZIP must contain exactly one .xlsx master list")
    return sources, masters[0]


def read_master(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data), data_only=False)
    try:
        candidates = []
        for sheet in workbook:
            headers = {}
            for cell in sheet[1]:
                label = str(cell.value or "").strip().casefold()
                for field, aliases in ALIASES.items():
                    if label in aliases:
                        if field in headers:
                            raise IntakeError(f"Ambiguous {field} columns in {sheet.title}")
                        headers[field] = cell.column
            if {"institution", "name", "address"} <= headers.keys():
                candidates.append((sheet, headers))
        if len(candidates) != 1:
            raise IntakeError("Expected exactly one master sheet with Institution, Supervisor and Email headers in row 1")
        sheet, headers = candidates[0]
        result = []
        for number in range(2, sheet.max_row + 1):
            raw_cells = {c.coordinate: c.value for c in sheet[number]}
            if all(value is None for value in raw_cells.values()):
                continue
            values, field_cells = {"profile": ""}, {}
            for field, column in headers.items():
                cell = sheet.cell(number, column)
                if field == "institution":
                    for merged in sheet.merged_cells.ranges:
                        if cell.coordinate in merged and merged.min_col == merged.max_col:
                            cell = sheet.cell(merged.min_row, column)
                            break
                if cell.data_type == "f":
                    raise IntakeError(f"Formula in identity field: {sheet.title}!{cell.coordinate}")
                values[field] = str(cell.value or "").strip()
                field_cells[field] = cell.coordinate
            if not values["institution"] or not values["name"]:
                raise IntakeError(f"Missing Institution or Supervisor at {sheet.title}!{number}; unmerged blanks are not filled")
            result.append({"sheet": sheet.title, "row": number, "values": values,
                           "evidence": {"values": values, "raw_cells": raw_cells, "field_cells": field_cells}})
        if not result:
            raise IntakeError("Master sheet contains no supervisor rows")
        return result
    finally:
        workbook.close()
